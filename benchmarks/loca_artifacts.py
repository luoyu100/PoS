"""LOCA-Bench bounded tool observation materialization."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any


_ERROR_PREFIXES = (
    "error",
    "permission denied",
    "traceback",
    "execution timed out",
    "tool execution failed",
)
_RESULT_MANIFEST_KEYS = {
    "path",
    "record_count",
    "schema",
    "validation_checks",
    "validation_status",
}


def materialize_loca_observation(
    observation: str,
    *,
    artifact_dir: Path,
    agent_artifact_dir: str,
    case_id: str,
    action: dict[str, Any],
    tool_use_counter: int,
    inline_max_chars: int,
    preview_items: int,
    structured_only: bool = True,
) -> str:
    """Persist a large LOCA tool result and return a compact agent view.

    The public adapter contract remains unchanged: both the input and return
    value are strings.  Results are materialized only when the observation can
    be decoded as LOCA's OpenAI tool-message list and the tool content exceeds
    ``inline_max_chars``.
    """

    if len(observation) <= inline_max_chars:
        return observation

    messages = _tool_messages(observation)
    if messages is None or len(messages) != 1:
        return observation

    message = messages[0]
    content = message.get("content")
    if not isinstance(content, str) or len(content) <= inline_max_chars:
        return observation
    if content.lstrip().lower().startswith(_ERROR_PREFIXES):
        return observation

    payload, payload_start = _structured_payload(content)
    if payload is None and structured_only:
        return observation
    if payload is None:
        payload = content
        payload_start = None

    tool_name = str(action.get("name", "unknown_tool"))
    filename = f"tool_result_{tool_use_counter:04d}.json"
    artifact_path = artifact_dir / filename
    agent_path = f"{agent_artifact_dir.rstrip('/')}/{filename}"
    raw_sha256 = hashlib.sha256(content.encode("utf-8")).hexdigest()

    payload_type = _payload_type(payload)
    artifact = {
        "version": 1,
        "case_id": case_id,
        "tool_name": tool_name,
        "tool_call_id": message.get("tool_call_id"),
        "tool_use_counter": tool_use_counter,
        "payload_type": payload_type,
        "payload": payload,
        "raw_content": content,
        "raw_content_sha256": raw_sha256,
    }
    artifact_dir.mkdir(parents=True, exist_ok=True)
    artifact_path.write_text(
        json.dumps(artifact, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    preview = _preview(payload, preview_items)
    fields = _fields(payload)
    size_label = _size_label(payload)
    header = content[:payload_start].strip() if payload_start is not None else ""
    compact_content = _compact_content(
        header=header,
        tool_name=tool_name,
        agent_path=agent_path,
        payload_type=payload_type,
        size_label=size_label,
        fields=fields,
        preview=preview,
    )

    compact_message = dict(message)
    compact_message["content"] = compact_content
    return json.dumps([compact_message], ensure_ascii=False)


def validated_workspace_output(
    observation: str,
    *,
    agent_workspace: Path,
    producer: str,
    artifact_contracts: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any] | None:
    """Validated workspace output."""

    messages = _tool_messages(observation)
    if messages is None or len(messages) != 1:
        return None
    content = messages[0].get("content")
    if not isinstance(content, str):
        return None
    manifest = _result_manifest(content)
    if manifest is None:
        return None
    path = manifest.get("path")
    count = manifest.get("record_count")
    checks = manifest.get("validation_checks")
    status = str(manifest.get("validation_status", "")).strip().casefold()
    if (
        not isinstance(path, str)
        or isinstance(count, bool)
        or not isinstance(count, int)
        or count < 0
        or not isinstance(checks, dict)
        or not checks
        or status not in {"passed", "failed"}
    ):
        return None

    normalized_checks: dict[str, dict[str, Any]] = {}
    for name, check in checks.items():
        if not isinstance(name, str) or not name or not isinstance(check, dict):
            return None
        observed = check.get("observed")
        expected = check.get("expected")
        basis = check.get("basis")
        passed = check.get("passed")
        if (
            not _manifest_value(observed)
            or not _manifest_value(expected)
            or not isinstance(basis, str)
            or not basis.strip()
            or not isinstance(passed, bool)
            or passed != (observed == expected)
        ):
            return None
        normalized_checks[name] = {
            "observed": observed,
            "expected": expected,
            "basis": basis.strip(),
            "passed": passed,
        }

    manifest_passed = all(check["passed"] for check in normalized_checks.values())
    if (status == "passed") != manifest_passed:
        return None

    root = agent_workspace.resolve()
    try:
        candidate = Path(path).resolve()
        candidate.relative_to(root)
        if not candidate.is_file():
            return None
        stat = candidate.stat()
    except (OSError, ValueError):
        return None
    on_disk_count = _file_record_count(candidate)
    if on_disk_count is not None and count != on_disk_count:
        return None
    schema = manifest.get("schema")
    if not isinstance(schema, (str, list, dict)) or not schema:
        return None

    contract = (artifact_contracts or {}).get(str(candidate))
    if isinstance(contract, dict) and contract.get("original_headers"):
        observed_headers = _excel_headers(
            candidate,
            sheet_name=contract.get("sheet_name"),
        )
        expected_headers = contract["original_headers"]
        if observed_headers is None:
            return None
        normalized_checks["original_headers"] = {
            "observed": observed_headers,
            "expected": expected_headers,
            "basis": "runner-captured original workbook headers",
            "passed": observed_headers == expected_headers,
        }

        schema = observed_headers

    all_passed = all(check["passed"] for check in normalized_checks.values())
    status = "passed" if all_passed else "failed"
    summary = {
        name: {
            "observed": check["observed"],
            "expected": check["expected"],
            "passed": check["passed"],
        }
        for name, check in normalized_checks.items()
    }
    basis = {name: check["basis"] for name, check in normalized_checks.items()}
    return {
        "path": str(candidate),
        "bytes": stat.st_size,
        "change": "validated",
        "producer": producer,
        "record_count": count,
        "schema": (
            schema
            if isinstance(schema, str)
            else json.dumps(schema, ensure_ascii=False, sort_keys=True)
        ),
        "validation_status": status,
        "validation_check_count": len(normalized_checks),
        "validation_summary": json.dumps(summary, ensure_ascii=False, sort_keys=True),
        "validation_basis": json.dumps(basis, ensure_ascii=False, sort_keys=True),
    }


def capture_excel_artifact_contract(
    action: dict[str, Any],
    *,
    agent_workspace: Path,
) -> dict[str, Any] | None:
    """Capture excel artifact contract."""

    name = str(action.get("name", ""))
    if not name.endswith("excel_read_data_from_excel"):
        return None
    arguments = action.get("arguments", {})
    if not isinstance(arguments, dict):
        return None
    raw_path = arguments.get("filepath") or arguments.get("path")
    if not isinstance(raw_path, str):
        return None

    root = agent_workspace.resolve()
    try:
        path = Path(raw_path).resolve()
        path.relative_to(root)
    except ValueError:
        return None
    if path.suffix.casefold() not in {".xlsx", ".xlsm"} or not path.is_file():
        return None

    sheet_name = arguments.get("sheet_name")
    headers = _excel_headers(path, sheet_name=sheet_name)
    if not headers:
        return None
    return {
        "path": str(path),
        "artifact_type": "excel_workbook",
        "sheet_name": sheet_name,
        "original_headers": headers,
        "preserve_headers": True,
        "captured_by": name,
    }


def _excel_headers(path: Path, *, sheet_name: object = None) -> list[Any] | None:
    """Excel headers."""

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = (
                workbook[str(sheet_name)]
                if isinstance(sheet_name, str) and sheet_name in workbook.sheetnames
                else workbook.active
            )
            row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
            return list(row)
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError):
        return None


def _manifest_value(value: Any) -> bool:
    """Return whether a manifest value is compact, explicit JSON evidence."""

    if isinstance(value, bool) or value is None:
        return True
    if isinstance(value, (int, float, str)):
        return value != ""
    if isinstance(value, list):
        return bool(value) and all(_manifest_value(item) for item in value)
    if isinstance(value, dict):
        return bool(value) and all(
            isinstance(key, str) and key and _manifest_value(item)
            for key, item in value.items()
        )
    return False


def _file_record_count(path: Path) -> int | None:
    """Return a deterministic record count for common durable result files."""

    suffix = path.suffix.casefold()
    try:
        if suffix == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as handle:
                rows = [row for row in csv.reader(handle) if any(row)]
            return max(0, len(rows) - 1)
        if suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(payload, list):
                return len(payload)
    except (OSError, UnicodeError, csv.Error, json.JSONDecodeError):
        return None
    return None


def append_workspace_output_manifest(
    observation: str,
    *,
    outputs: list[dict[str, Any]],
) -> str:
    """Append runner-observed durable files to one LOCA tool observation.

    The manifest contains paths and metadata only; the output content remains
    in the agent workspace. It gives the next decision an exact, observed
    reopen handle without replaying a large result through the model context.
    """

    if not outputs:
        return observation
    messages = _tool_messages(observation)
    if messages is None or len(messages) != 1:
        return observation
    content = messages[0].get("content")
    if not isinstance(content, str):
        return observation

    manifest = json.dumps(outputs, ensure_ascii=False, sort_keys=True)
    compact_message = dict(messages[0])
    compact_message["content"] = (
        content.rstrip()
        + chr(10) * 2
        + "PoS durable workspace outputs (runner-observed):"
        + chr(10)
        + manifest
    )
    return json.dumps([compact_message], ensure_ascii=False)


def _tool_messages(observation: str) -> list[dict[str, Any]] | None:
    """Decode LOCA's serialized OpenAI tool-message observation."""

    try:
        messages = json.loads(observation)
    except (json.JSONDecodeError, TypeError):
        return None
    if not isinstance(messages, list):
        return None
    if not all(isinstance(message, dict) for message in messages):
        return None
    return messages


def _result_manifest(content: str) -> dict[str, Any] | None:
    """Extract a result manifest from Python stdout before wrapper metadata."""

    stdout = content
    if "=== STDOUT ===" in stdout:
        stdout = stdout.split("=== STDOUT ===", 1)[1]
    if "=== EXECUTION INFO ===" in stdout:
        stdout = stdout.split("=== EXECUTION INFO ===", 1)[0]

    decoder = json.JSONDecoder()
    candidates = [match.start(1) for match in re.finditer(r"(?m)^[ \t]*({)", stdout)]
    for start in reversed(candidates):
        try:
            payload, _ = decoder.raw_decode(stdout, start)
        except json.JSONDecodeError:
            continue
        if isinstance(payload, dict) and _RESULT_MANIFEST_KEYS.issubset(payload):
            return payload
    return None


def _structured_payload(content: str) -> tuple[Any | None, int | None]:
    """Return the trailing JSON array/object embedded in a tool response."""

    decoder = json.JSONDecoder()
    candidates = [
        match.start(1) for match in re.finditer(r"(?m)^[ \t]*([\[{])", content)
    ]
    for start in reversed(candidates):
        try:
            payload, end = decoder.raw_decode(content, start)
        except json.JSONDecodeError:
            continue
        if content[end:].strip():
            continue
        if isinstance(payload, (list, dict)):
            return payload, start
    return None, None


def _payload_type(payload: Any) -> str:
    """Describe the representation available under artifact['payload']."""

    if isinstance(payload, list):
        return "json_array"
    if isinstance(payload, dict):
        return "json_object"
    return "raw_text"


def _preview(payload: Any, preview_items: int) -> Any:
    """Build a bounded schema preview without changing exact payload values."""

    limit = max(0, preview_items)
    if isinstance(payload, list):
        return payload[:limit]
    if isinstance(payload, dict):
        return dict(list(payload.items())[:limit])
    if isinstance(payload, str):
        head_chars = max(200, min(1200, 300 * max(1, limit)))
        if len(payload) <= head_chars * 2:
            return payload
        return {
            "head": payload[:head_chars],
            "tail": payload[-head_chars:],
            "omitted_chars": len(payload) - 2 * head_chars,
        }
    return None


def _fields(payload: Any) -> list[str]:
    """Return stable top-level fields for a tabular or object payload."""

    if isinstance(payload, list):
        fields: list[str] = []
        for item in payload[:10]:
            if not isinstance(item, dict):
                continue
            for key in item:
                if key not in fields:
                    fields.append(str(key))
        return fields
    if isinstance(payload, dict):
        return [str(key) for key in payload]
    return []


def _size_label(payload: Any) -> str:
    """Describe the number of rows/items without claiming a table schema."""

    if isinstance(payload, list):
        return f"Rows: {len(payload)}"
    if isinstance(payload, dict):
        return f"Top-level items: {len(payload)}"
    if isinstance(payload, str):
        return f"Raw text characters: {len(payload)}"
    return "Payload size: unavailable"


def _compact_content(
    *,
    header: str,
    tool_name: str,
    agent_path: str,
    payload_type: str,
    size_label: str,
    fields: list[str],
    preview: Any,
) -> str:
    """Render the bounded observation shown to the task agent."""

    parts = []
    if header:
        parts.append(header)
    parts.extend(
        [
            "The complete tool result was materialized because it is "
            "too large to reproduce safely inside another tool call.",
            f"Tool: {tool_name}",
            size_label,
            f"Payload type: {payload_type}",
            f"Fields: {', '.join(fields) if fields else 'not tabular'}",
            f"Artifact: {agent_path}",
            "Canonical Python access: artifact = json.load(open(artifact_path, "
            "encoding='utf-8')); payload = artifact['payload'].",
            "Interpret payload using the Payload type and Fields above. For "
            "raw_text, payload is the exact original tool content. The "
            "artifact path is intended for Python; filesystem tools may have "
            "narrower directory access.",
            "Preview (for schema inspection only):",
            json.dumps(preview, ensure_ascii=False, indent=2),
            "Use the complete artifact payload for calculations. Do not copy "
            "the complete payload into action arguments or Python source code.",
        ]
    )
    return "\n\n".join(parts)
